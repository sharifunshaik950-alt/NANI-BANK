from database import connect_database

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def statement(account_number):

    print("\n========== STATEMENT ==========")

    connection = connect_database()
    cursor = connection.cursor()

    cursor.execute("""
        SELECT customer_name, phone, address, balance
        FROM customers
        WHERE account_number = %s
    """, (account_number,))

    customer = cursor.fetchone()

    if customer is None:

        print("Account not found.")

        cursor.close()
        connection.close()
        return

    name = customer[0]
    phone = customer[1]
    address = customer[2]
    balance = customer[3]

    cursor.execute("""
        SELECT
            transaction_date,
            transaction_type,
            amount,
            balance_after
        FROM transactions
        WHERE account_number = %s
        ORDER BY transaction_date
    """, (account_number,))

    transactions = cursor.fetchall()

    cursor.close()
    connection.close()

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Statement"

    # Bank title
    sheet["A1"] = "NANI BANK"
    sheet["A1"].font = Font(size=18, bold=True)
    sheet["A1"].alignment = Alignment(horizontal="center")

    sheet.merge_cells("A1:D1")

    # Customer details
    sheet["A3"] = "Customer Name"
    sheet["B3"] = name

    sheet["A4"] = "Account Number"
    sheet["B4"] = account_number

    sheet["A5"] = "Phone"
    sheet["B5"] = phone

    sheet["A6"] = "Address"
    sheet["B6"] = address

    # Table headings
    headers = [
        "Date",
        "Transaction",
        "Amount",
        "Balance"
    ]

    for column, header in enumerate(headers, start=1):

        cell = sheet.cell(
            row=8,
            column=column
        )

        cell.value = header
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")

    # Transaction data
    row = 9

    for transaction in transactions:

        date = transaction[0]
        transaction_type = transaction[1]
        amount = transaction[2]
        balance_after = transaction[3]

        sheet.cell(row=row, column=1).value = date
        sheet.cell(row=row, column=2).value = transaction_type
        sheet.cell(row=row, column=3).value = float(amount)
        sheet.cell(row=row, column=4).value = float(balance_after)

        row += 1

    # Final balance
    sheet.cell(row=row + 1, column=3).value = "Final Balance"
    sheet.cell(row=row + 1, column=3).font = Font(bold=True)

    sheet.cell(row=row + 1, column=4).value = float(balance)
    sheet.cell(row=row + 1, column=4).font = Font(bold=True)

    # Borders
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin")
    )

    for row_cells in sheet.iter_rows(
        min_row=8,
        max_row=row - 1,
        min_col=1,
        max_col=4
    ):

        for cell in row_cells:
            cell.border = thin_border

    # Column widths
    for column in range(1, 5):

        letter = get_column_letter(column)

        sheet.column_dimensions[letter].width = 22

    # Excel file name
    file_name = f"statement_{account_number}.xlsx"

    workbook.save(file_name)

    print("\nStatement created successfully!")
    print("Excel file:", file_name)
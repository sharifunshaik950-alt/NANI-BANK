from database import connect_database

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side


# =========================================================
# 1. SEARCH CUSTOMER
# =========================================================

def search_customer():

    print("\n========== SEARCH CUSTOMER ==========")

    search = input("Enter customer name or account number: ")

    connection = connect_database()
    cursor = connection.cursor()

    if search.isdigit():

        cursor.execute("""
            SELECT
                account_number,
                customer_name,
                phone,
                address,
                balance
            FROM customers
            WHERE account_number = %s
        """, (int(search),))

    else:

        cursor.execute("""
            SELECT
                account_number,
                customer_name,
                phone,
                address,
                balance
            FROM customers
            WHERE customer_name ILIKE %s
        """, ('%' + search + '%',))

    customers = cursor.fetchall()

    cursor.close()
    connection.close()

    if not customers:

        print("Customer not found.")
        return

    for customer in customers:

        print("\n----------------------------")
        print("Account Number:", customer[0])
        print("Customer Name :", customer[1])
        print("Phone         :", customer[2])
        print("Address       :", customer[3])
        print("Balance       :", customer[4])


# =========================================================
# 2. CUSTOMERS LIST
# =========================================================

def customers_list():

    print("\n========== CUSTOMERS LIST ==========")

    connection = connect_database()
    cursor = connection.cursor()

    cursor.execute("""
        SELECT
            account_number,
            customer_name,
            phone,
            balance
        FROM customers
        ORDER BY customer_name
    """)

    customers = cursor.fetchall()

    cursor.close()
    connection.close()

    if not customers:

        print("No customers found.")
        return

    print("\n------------------------------------------------------------")
    print("Account Number     Name                 Phone        Balance")
    print("------------------------------------------------------------")

    for customer in customers:

        print(
            f"{customer[0]}    "
            f"{customer[1]:20} "
            f"{customer[2]:12} "
            f"{customer[3]}"
        )


# =========================================================
# 3. ALL CUSTOMERS EXCEL
# =========================================================

def all_customers_excel():

    print("\n========== ALL CUSTOMERS EXCEL ==========")

    connection = connect_database()
    cursor = connection.cursor()

    # Get all customers
    cursor.execute("""
        SELECT
            account_number,
            customer_name,
            phone,
            address,
            balance
        FROM customers
        ORDER BY customer_name
    """)

    customers = cursor.fetchall()

    if not customers:

        print("No customers found.")

        cursor.close()
        connection.close()

        return

    # Create Excel workbook
    workbook = Workbook()

    sheet = workbook.active
    sheet.title = "All Customers"

    # =====================================================
    # TITLE
    # =====================================================

    sheet["A1"] = "NANI BANK - ALL CUSTOMERS"

    sheet["A1"].font = Font(
        size=18,
        bold=True
    )

    sheet["A1"].alignment = Alignment(
        horizontal="center"
    )

    sheet.merge_cells("A1:H1")

    # =====================================================
    # HEADINGS
    # =====================================================

    headers = [
        "Account Number",
        "Customer Name",
        "Phone",
        "Address",
        "Transaction Date",
        "Transaction",
        "Amount",
        "Balance"
    ]

    for column, header in enumerate(headers, start=1):

        cell = sheet.cell(
            row=3,
            column=column
        )

        cell.value = header

        cell.font = Font(
            bold=True
        )

        cell.alignment = Alignment(
            horizontal="center"
        )

    # Starting Excel row
    row = 4

    # =====================================================
    # ADD ALL CUSTOMERS
    # =====================================================

    for customer in customers:

        account_number = customer[0]
        customer_name = customer[1]
        phone = customer[2]
        address = customer[3]
        current_balance = customer[4]

        # Get transactions for this customer
        cursor.execute("""
            SELECT
                transaction_date,
                transaction_type,
                amount,
                balance_after
            FROM transactions
            WHERE account_number = %s
            ORDER BY transaction_date
        """, (account_number,))

        transactions = cursor.fetchall()

        # =================================================
        # CUSTOMER HAS TRANSACTIONS
        # =================================================

        if transactions:

            for transaction in transactions:

                sheet.cell(
                    row=row,
                    column=1
                ).value = account_number

                sheet.cell(
                    row=row,
                    column=2
                ).value = customer_name

                sheet.cell(
                    row=row,
                    column=3
                ).value = phone

                sheet.cell(
                    row=row,
                    column=4
                ).value = address

                sheet.cell(
                    row=row,
                    column=5
                ).value = transaction[0]

                sheet.cell(
                    row=row,
                    column=6
                ).value = transaction[1]

                sheet.cell(
                    row=row,
                    column=7
                ).value = float(transaction[2])

                sheet.cell(
                    row=row,
                    column=8
                ).value = float(transaction[3])

                row += 1

        # =================================================
        # CUSTOMER HAS NO TRANSACTIONS
        # =================================================

        else:

            sheet.cell(
                row=row,
                column=1
            ).value = account_number

            sheet.cell(
                row=row,
                column=2
            ).value = customer_name

            sheet.cell(
                row=row,
                column=3
            ).value = phone

            sheet.cell(
                row=row,
                column=4
            ).value = address

            sheet.cell(
                row=row,
                column=8
            ).value = float(current_balance)

            row += 1

    # =====================================================
    # BORDERS
    # =====================================================

    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin")
    )

    for row_cells in sheet.iter_rows(
        min_row=3,
        max_row=row - 1,
        min_col=1,
        max_col=8
    ):

        for cell in row_cells:

            cell.border = thin_border

    # =====================================================
    # COLUMN WIDTHS
    # =====================================================

    widths = {
        "A": 20,
        "B": 20,
        "C": 15,
        "D": 20,
        "E": 25,
        "F": 15,
        "G": 15,
        "H": 15
    }

    for column, width in widths.items():

        sheet.column_dimensions[column].width = width

    # =====================================================
    # SAVE EXCEL FILE
    # =====================================================

    file_name = "all_customers.xlsx"

    workbook.save(file_name)

    cursor.close()
    connection.close()

    print("\nExcel file created successfully!")
    print("File name:", file_name)